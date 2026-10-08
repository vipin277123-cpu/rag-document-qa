from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class DocumentPage:
    text: str
    page_number: int | None
    metadata: dict[str, str] = field(default_factory=dict)


class DocumentProcessor:
    """Extract readable text while retaining page and document properties."""

    def extract(self, file_path: Path) -> list[DocumentPage]:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(file_path)
        if suffix == ".txt":
            return self._extract_txt(file_path)
        if suffix == ".docx":
            return self._extract_docx(file_path)
        raise ValueError("Unsupported file type. Upload a PDF, TXT, or DOCX document.")

    @staticmethod
    def _extract_pdf(file_path: Path) -> list[DocumentPage]:
        reader = PdfReader(str(file_path))
        pages = [
            DocumentPage(text=(page.extract_text() or "").strip(), page_number=index)
            for index, page in enumerate(reader.pages, start=1)
        ]
        if not any(page.text for page in pages):
            raise ValueError("No extractable text was found in the PDF.")
        return pages

    @staticmethod
    def _extract_txt(file_path: Path) -> list[DocumentPage]:
        text = file_path.read_text(encoding="utf-8-sig").strip()
        if not text:
            raise ValueError("The document is empty.")
        return [DocumentPage(text=text, page_number=1)]

    @staticmethod
    def _extract_docx(file_path: Path) -> list[DocumentPage]:
        document = Document(str(file_path))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text = "\n\n".join(blocks).strip()
        if not text:
            raise ValueError("No extractable text was found in the DOCX document.")

        properties = document.core_properties
        metadata = {
            key: value.strip()
            for key, value in {
                "title": properties.title,
                "author": properties.author,
                "subject": properties.subject,
            }.items()
            if value and value.strip()
        }
        return [DocumentPage(text=text, page_number=1, metadata=metadata)]
